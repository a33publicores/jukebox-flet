"""Convierte los pedidos a Excel (.xlsx). Lo usan la API (panel web) y el reproductor."""
import io
import time

COLUMNAS = [
    ("Fecha y hora", "fecha"), ("Usuario", "usuario"), ("Canción", "titulo"),
    ("Canal", "canal"), ("videoId", "video_id"), ("Estado", "estado"),
    ("Estado2", "estado2"), ("Orden", "orden"), ("Id", "id"),
]
ZONA_HORAS = -5


def fecha_local(ts):
    try:
        return time.strftime("%Y-%m-%d %H:%M:%S", time.gmtime(float(ts) + ZONA_HORAS * 3600))
    except (TypeError, ValueError):
        return ""


def excel_bytes(filas, titulo="Pedidos"):
    """filas: lista de dicts como los de la tabla pedidos."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = str(titulo)[:31] or "Pedidos"
    ws.append([c for c, _ in COLUMNAS])
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="0B1220")
    for f in filas:
        fila = []
        for _, k in COLUMNAS:
            v = fecha_local(f.get("ts")) if k == "fecha" else f.get(k)
            fila.append("" if v is None else v)
        ws.append(fila)
    for col, ancho in zip("ABCDEFGHI", (20, 14, 60, 26, 14, 11, 16, 8, 8)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
